"""Tests for building decks as native slides on a base deck's master."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_sg import make_pptx, run, write  # noqa: E402


def test_build_deck(tmp_path):
    from pptx import Presentation
    base = make_pptx(tmp_path / "base.pptx", ["Intro", "Old topic", "Scalability", "Hidden extra"], hidden={4})
    spec = write(tmp_path / "deck.json", {"base": "base.pptx", "output": "out/lecture.pptx", "slides": [
        {"copy": 3, "edits": [{"find": "Bullet about Scalability", "replace": "Scalability at every level"}],
         "insert": [{"after": "Scalability at every level", "lines": ["Inserted line"]}],
         "notes": "added note", "notes_mode": "append"},
        {"layout": "Title Only", "title": "Hash partitioning", "notes": "speaker text", "shapes": [
            {"type": "box", "x": 1, "y": 1.5, "w": 2, "h": 0.6, "text": "key: **customer_id**"},
            {"type": "arrow", "from": [3, 1.8], "to": [5, 1.8], "label": "hash mod 3"},
            {"type": "table", "x": 1, "y": 3, "w": 6, "h": 1, "rows": [["", "Hash", "Range"], ["Even", "yes", "no"]]}]},
        {"copy": 1},
        {"layout": "Title and Content", "title": "Levels", "bullets": ["**Infrastructure**", {"text": "compute", "level": 1}]}]})
    r = run("pptx", "build", "--spec", spec)
    assert r["ok"] and r["slides"] == 4, r
    prs = Presentation(str(tmp_path / "out" / "lecture.pptx"))
    titles = [s.shapes.title.text for s in prs.slides]
    assert titles == ["Scalability", "Hash partitioning", "Intro", "Levels"]   # spec order; others dropped
    first = prs.slides[0]
    assert first.placeholders[1].text_frame.text.split("\n")[:3] == ["Scalability at every level", "Inserted line",
                                                                   "Second line"]
    assert first.notes_slide.notes_text_frame.text == "Original note 3\n\nadded note"
    shapes = prs.slides[1].shapes
    assert any(s.has_text_frame and "customer_id" in s.text_frame.text for s in shapes)
    assert any(s.has_table for s in shapes)
    bold = [r for s in prs.slides[3].placeholders if s.placeholder_format.idx == 1
            for p in s.text_frame.paragraphs for r in p.runs if r.font.bold]
    assert [r.text for r in bold] == ["Infrastructure"]
    bad = write(tmp_path / "bad.json", {"base": "base.pptx", "output": "x.pptx", "slides": [
        {"copy": 2, "edits": [{"find": "not on the slide", "replace": "x"}]}, {"copy": 9}]})
    r = run("pptx", "build", "--spec", bad)
    assert not r["ok"] and "outside the base deck" in r["error"]


def test_spec_example_builds(tmp_path):
    """The example generator shipped with the plugin writes a specification that builds."""
    import subprocess
    import sys as _sys
    base = make_pptx(tmp_path / "template.pptx", ["Course title"])
    example = Path(__file__).resolve().parents[1] / "references" / "deck_spec_example.py"
    subprocess.run([_sys.executable, str(example), "--base", str(base), "--out", str(tmp_path / "L01" / "deck.json")],
                   check=True, capture_output=True)
    r = run("pptx", "build", "--spec", tmp_path / "L01" / "deck.json")
    assert r["ok"] and r["slides"] == 5 and not r["problems"], r


def test_edit_table_cell(tmp_path):
    """Edits on a copied slide also find text in table cells."""
    from pptx import Presentation
    from pptx.util import Inches
    base = make_pptx(tmp_path / "base.pptx", ["Intro", "Governance"])
    prs = Presentation(str(base))
    table = prs.slides[1].shapes.add_table(2, 2, Inches(1), Inches(3), Inches(6), Inches(1)).table
    table.cell(0, 0).text, table.cell(1, 0).text = "Governance", "Sets the rules once a year"
    prs.save(str(base))
    spec = write(tmp_path / "deck.json", {"base": "base.pptx", "output": "out.pptx", "slides": [
        {"copy": 2, "edits": [{"find": "once a year", "replace": "for every table"}]}]})
    r = run("pptx", "build", "--spec", spec)
    assert r["ok"], r
    cells = [s for s in Presentation(str(tmp_path / "out.pptx")).slides[0].shapes if s.has_table][0].table
    assert cells.cell(1, 0).text == "Sets the rules for every table"


def test_literal_asterisk(tmp_path):
    """\\* in slide text is a literal asterisk, not italic markup."""
    from pptx import Presentation
    make_pptx(tmp_path / "base.pptx", ["Intro"])
    spec = write(tmp_path / "deck.json", {"base": "base.pptx", "output": "out.pptx", "slides": [
        {"layout": "Title and Content", "title": "Paths", "bullets": ["read orders/\\*/\\*.parquet and *this*"]}]})
    assert run("pptx", "build", "--spec", spec)["ok"]
    para = Presentation(str(tmp_path / "out.pptx")).slides[0].placeholders[1].text_frame.paragraphs[0]
    assert para.text == "read orders/*/*.parquet and this" and para.runs[-1].font.italic
