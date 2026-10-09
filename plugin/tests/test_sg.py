"""Tests for the sg tool on a synthetic two-deck course.

Run: python3 -m pytest plugin/tests
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest
import yaml

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import sg  # noqa: E402
from sglib.findings import _priority  # noqa: E402
from sglib.sources import match_quote  # noqa: E402

SOURCE_TEXT = """Apache Parquet is a column-oriented data file format designed for efficient
data storage and retrieval. It provides high performance compression and encoding
schemes to handle complex data in bulk. Row groups are a logical horizontal
partitioning of the data into rows. Since version 2.5 the format defines a page
index that stores min and max statistics for every page.""" * 3


def run(*argv) -> dict:
    """Run sg in-process and return its JSON summary."""
    import contextlib
    import io
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        code = sg.main([str(a) for a in argv])
    out = json.loads(buf.getvalue())
    out["_code"] = code
    return out


def write(path: Path, data) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data) if not isinstance(data, str) else data, encoding="utf-8")
    return path


def make_pptx(path: Path, titles: list[str], hidden: set[int] = frozenset()) -> Path:
    from pptx import Presentation
    prs = Presentation()
    for i, title in enumerate(titles, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = f"Bullet about {title}\nSecond line"
        slide.notes_slide.notes_text_frame.text = f"Original note {i}"
        if i in hidden:
            slide._element.set("show", "0")
    prs.save(str(path))
    return path


# ---------------------------------------------------------------- unit tests

def test_match_quote_normalizes_quotes_dashes_and_whitespace():
    text = "The format “stores” data—column by column\n  in row groups of 128 MB."
    ok, start, end = match_quote('the format "stores" data-column by column in row groups', text)
    assert ok and text[start:end].startswith("The format")


def test_match_quote_ellipsis_parts_must_appear_in_order():
    text = "Row groups split rows horizontally. Column chunks hold one column. Pages hold values."
    assert match_quote("Row groups split rows horizontally ... Pages hold values", text)[0]
    assert not match_quote("Pages hold values ... Row groups split rows horizontally", text)[0]


def test_match_quote_rejects_short_and_paraphrased_quotes():
    text = "Parquet uses dictionary encoding when values repeat often."
    assert not match_quote("Parquet", text)[0]
    assert not match_quote("Parquet applies dictionary encoding when values repeat", text)[0]


def _finding(**kw):
    base = {"maturity": "ga", "objectives": ["file formats"], "invalidates": [],
            "teaching_signals": [{"vendor": False}, {"vendor": False}]}
    return {**base, **kw}


def test_priority_counts_only_independent_signals():
    vendor_only = _finding(teaching_signals=[{"vendor": True}, {"vendor": True}])
    assert _priority(vendor_only, set())[0] == "SHOULD"


def test_priority_rules_and_maturity_caps():
    assert _priority(_finding(), set())[0] == "MUST"
    assert _priority(_finding(teaching_signals=[]), set())[0] == "SHOULD"
    assert _priority(_finding(objectives=[], teaching_signals=[]), set())[0] == "OPTIONAL"
    assert _priority(_finding(maturity="preview"), set())[0] == "SHOULD"
    assert _priority(_finding(maturity="announced", invalidates=["c1"]), {"c1"})[0] == "OPTIONAL"
    assert _priority(_finding(objectives=[], teaching_signals=[], invalidates=["c1"]), {"c1"})[0] == "MUST"


# ---------------------------------------------------------------- end-to-end

@pytest.fixture()
def course(tmp_path: Path) -> dict:
    root = tmp_path / "course-set"
    deck_a = make_pptx(tmp_path / "a.pptx", ["Formats", "Parquet layout", "Parquet stats", "Old slide"], hidden={4})
    deck_b = make_pptx(tmp_path / "b.pptx", ["Engines", "Streaming preview"])
    manifest = {"name": "test", "decks": [
        {"id": "deck-a", "path": str(deck_a), "minutes": 90, "last_revised": 2022,
         "recreate": [{"slides": "2-3", "topic": "Parquet"}]},
        {"id": "deck-b", "path": str(deck_b), "minutes": 90, "last_revised": 2025,
         "out_of_scope": [{"slides": "2", "reason": "separate lecture"}]}]}
    m = write(root / "course.yaml", yaml.safe_dump(manifest))
    source = write(tmp_path / "parquet-docs.txt", SOURCE_TEXT)
    return {"root": root, "manifest": m, "source": source, "decks": {"deck-a": deck_a, "deck-b": deck_b}}


def test_full_report_flow(course):
    root, m = course["root"], course["manifest"]
    store = root / "sources"
    ws = {d: root / d / "workspace" for d in course["decks"]}

    # ingest
    for d, path in course["decks"].items():
        r = run("ingest", path, "--out", ws[d])
        assert r["_code"] == 0, r
    slides_a = json.loads((ws["deck-a"] / "slides.json").read_text())
    assert slides_a["slides"][3]["hidden"] and slides_a["slides"][1]["title"] == "Parquet layout"

    # fetch a local source
    r = run("fetch", "--file", course["source"], "--store", store)
    assert r["ok"], r
    fid = r["fetched"]

    # claims: one extraction part per deck
    write(ws["deck-a"] / "claims" / "extract_a.json", {"deck": "deck-a", "part": "a", "slides": [1, 4], "claims": [
        {"id": "claim:deck-a:a-001", "slide": 2, "slide_title": "Parquet layout", "text": "Row groups partition rows",
         "statement": "Parquet row groups partition data horizontally.", "kind": ["named"], "from_image": False},
        {"id": "claim:deck-a:a-002", "slide": 3, "slide_title": "Parquet stats", "text": "Stats only per row group",
         "statement": "Parquet keeps statistics only per row group.", "kind": ["named"], "from_image": False},
        {"id": "claim:deck-a:a-003", "slide": 4, "slide_title": "Old slide", "text": "hidden", "statement": "x",
         "kind": ["general"], "from_image": False}]})
    write(ws["deck-b"] / "claims" / "extract_a.json", {"deck": "deck-b", "part": "a", "slides": [1, 2], "claims": [
        {"id": "claim:deck-b:a-001", "slide": 1, "slide_title": "Engines", "text": "Engines read Parquet",
         "statement": "Query engines read Parquet files.", "kind": ["general"], "from_image": False},
        {"id": "claim:deck-b:a-002", "slide": 2, "slide_title": "Streaming preview", "text": "out of scope",
         "statement": "y", "kind": ["general"], "from_image": False}]})
    for d in ws:
        r = run("claims", "merge", ws[d], "--manifest", m)
        assert r["ok"], r
        assert len(r["dropped_excluded_slides"]) == 1  # hidden slide in A, out-of-scope slide in B
        assert run("claims", "select", ws[d], "--batch-size", 6)["ok"]

    src = {"url": "file:parquet-docs.txt", "title": "Parquet docs", "published": "2024",
           "retrieved": "2026-10-08", "fetched": fid, "role": "supports",
           "quote": "Row groups are a logical horizontal partitioning of the data into rows."}
    later = {**src, "role": "later_state",
             "quote": "the format defines a page index that stores min and max statistics for every page"}
    fake = {**src, "quote": "Parquet was invented in 1975 by a committee."}
    write(ws["deck-a"] / "research" / "batch_b01.json", {"batch": "b01", "claims": [
        {"claim_id": "claim:deck-a:a-001", "searched": ["parquet row groups"], "sources": [src, fake]},
        {"claim_id": "claim:deck-a:a-002", "searched": ["parquet page index"], "sources": [src, later]}]})
    write(ws["deck-b"] / "research" / "batch_b01.json", {"batch": "b01", "claims": [
        {"claim_id": "claim:deck-b:a-1", "searched": ["engines"], "sources": [{**src, "fetched": None}]}]})
    r = run("quotes", ws["deck-b"], "--store", store)   # a mistyped claim id must not detach the evidence
    assert not r["ok"] and "claim:deck-b:a-1" in r["error"], r
    write(ws["deck-b"] / "research" / "batch_b01.json", {"batch": "b01", "claims": [
        {"claim_id": "claim:deck-b:a-001", "searched": ["engines"], "sources": [{**src, "fetched": None}]}]})
    r = run("quotes", ws["deck-a"], "--store", store)
    assert r["matched"] == 3 and r["not_found"] == ["claim:deck-a:a-001#1"], r
    assert run("quotes", ws["deck-b"], "--store", store)["not_machine_checked"] == 1

    verified = {"verdict": "VERIFIED", "severity": None, "fix_kind": None, "outdated": None,
                "issue": None, "suggested_fix": None, "sources_used": [0]}
    outdated = {"verdict": "OUTDATED", "severity": "CRITICAL", "fix_kind": "new-content",
                "outdated": {"label": "changed-since", "change_date": "2018", "supporting_source": 0,
                             "changing_source": 1, "beyond_target_versions": False},
                "issue": "Page index added", "suggested_fix": "Mention page index", "sources_used": [0, 1]}
    write(ws["deck-a"] / "checks" / "b01.json", {"batch": "b01", "verdicts": [
        {"claim_id": "claim:deck-a:a-001", **verified}, {"claim_id": "claim:deck-a:a-002", **outdated}]})
    write(ws["deck-b"] / "checks" / "b01.json", {"batch": "b01", "verdicts": [
        {"claim_id": "claim:deck-b:a-001", **verified}]})

    # A CRITICAL verdict needs two second checks before the summary is complete.
    r = run("claims", "summary", ws["deck-a"])
    assert r["_code"] == 1 and r["needs_second_check"] == ["claim:deck-a:a-002"], r
    for k, sev in (("s1", "CRITICAL"), ("s2", "MAJOR")):
        write(ws["deck-a"] / "checks" / f"second_{k}.json", {"batch": "b01", "second_check": True, "verdicts": [
            {"claim_id": "claim:deck-a:a-002", **outdated, "severity": sev}]})
    r = run("claims", "summary", ws["deck-a"])
    assert r["ok"] and r["by_verdict"] == {"VERIFIED": 1, "OUTDATED": 1}, r
    row = [c for c in json.loads((ws["deck-a"] / "claim_check.json").read_text())["claims"]
           if c["claim_id"] == "claim:deck-a:a-002"][0]
    assert row["severity"] == "CRITICAL" and row["second_checked"]  # 2 of 3 votes
    assert run("claims", "summary", ws["deck-b"])["ok"]

    # topic maps, clusters, findings
    for d, sections in (("deck-a", [("S1", "Formats", [1, 3])]), ("deck-b", [("S1", "Engines", [1, 1])])):
        write(ws[d] / "topic_map.json", {"deck": d, "sections": [
            {"id": sid, "title": t, "slides": rng, "subtopics": ["x"], "named": [{"name": "Parquet", "kind": "standard"}],
             "objectives": ["file formats"], "content_date": "2019", "content_date_basis": "examples"}
            for sid, t, rng in sections]})
    write(root / "course" / "clusters.json", {"clusters": [
        {"id": "C1", "title": "Formats", "scope": "file formats", "sections": [{"deck": "deck-a", "section": "S1"}]}]})
    r = run("findings", "clusters", "--manifest", m)
    assert not r["ok"] and "deck-b/S1" in r["error"]
    write(root / "course" / "clusters.json", {"clusters": [
        {"id": "C1", "title": "Formats", "scope": "file formats", "sections": [{"deck": "deck-a", "section": "S1"}]},
        {"id": "C2", "title": "Engines", "scope": "engines", "sections": [{"deck": "deck-b", "section": "S1"}]}]})
    assert run("findings", "clusters", "--manifest", m)["ok"]

    signal = {"url": "https://example.org/survey", "title": "Survey", "kind": "survey", "vendor": False}
    write(root / "course" / "findings" / "cluster_C1.json", {"cluster": "C1", "searched": ["parquet 2025"], "findings": [
        {"key": "parquet-page-index", "statement": "Parquet page index", "date": "2019", "primary_source": src,
         "maturity": "ga", "sections": [{"deck": "deck-a", "section": "S1"}], "objectives": ["file formats"],
         "invalidates": ["claim:deck-a:a-002"], "teaching_signals": [signal], "relevance": "stats"},
        {"key": "hyped-format", "statement": "New format announced", "date": "2026", "primary_source": src,
         "maturity": "announced", "sections": [{"deck": "deck-a", "section": "S1"}], "objectives": ["file formats"],
         "invalidates": [], "teaching_signals": [], "relevance": "maybe"}]})
    r = run("findings", "merge", "--manifest", m)
    assert not r["ok"] and "review" in r["error"]
    write(root / "course" / "findings" / "review_C1.json", {"cluster": "C1", "reviews": [
        {"key": "parquet-page-index", "decision": "keep", "reason": "checked", "corrections": {}, "already_covered": None},
        {"key": "hyped-format", "decision": "keep", "reason": "real", "corrections": {}, "already_covered": None}]})
    r = run("findings", "merge", "--manifest", m)
    assert r["ok"] and r["by_priority"] == {"MUST": 1, "OPTIONAL": 1}, r
    merged = json.loads((root / "course" / "findings.json").read_text())
    must = merged["findings"][0]
    assert must["key"] == "parquet-page-index" and must["decks"]["deck-a"]["age"] == "missing"

    # plans
    write(ws["deck-a"] / "plan.json", {"deck": "deck-a", "items": [
        {"id": "P001", "type": "EDIT", "slides": [1], "current_text": "x", "proposed": "y", "reason": "r",
         "claims": [], "findings": [], "slides_delta": 0}]})
    r = run("plan", "validate", "--manifest", m, "--deck", "deck-a")
    assert not r["ok"] and "no REBUILD" in r["error"] and "parquet-page-index" in r["error"], r
    write(ws["deck-a"] / "plan.json", {"deck": "deck-a", "items": [
        {"id": "P001", "type": "REBUILD", "slides": [2, 3], "current_text": None, "proposed": "Outline",
         "reason": "Recreate Parquet", "claims": ["claim:deck-a:a-002"], "findings": ["parquet-page-index"],
         "slides_delta": 1},
        {"id": "P002", "type": "REMOVE", "slides": [4], "current_text": None, "proposed": "remove",
         "reason": "hidden", "claims": [], "findings": [], "slides_delta": -1}]})
    r = run("plan", "validate", "--manifest", m, "--deck", "deck-a")
    assert not r["ok"] and "excluded (hidden)" in r["error"], r
    plan = json.loads((ws["deck-a"] / "plan.json").read_text())
    plan["items"] = plan["items"][:1]
    write(ws["deck-a"] / "plan.json", plan)
    r = run("plan", "validate", "--manifest", m, "--deck", "deck-a")
    assert r["ok"] and r["rating"] == "significant update" and r["budget"]["net_slides_delta"] == 1, r
    write(ws["deck-b"] / "plan.json", {"deck": "deck-b", "items": []})
    assert run("plan", "validate", "--manifest", m, "--deck", "deck-b")["rating"] == "light refresh"

    # rendered outputs
    r = run("plan", "report", "--manifest", m, "--deck", "deck-a")
    assert r["ok"], r
    report = (root / "deck-a" / "output" / "modernization_report.md").read_text()
    assert "| Outdated | 1 |" in report
    assert "| Quotes matched / not found / too short / not machine-checked | 3 / 1 / 0 / 0 |" in report
    assert "QUOTE NOT FOUND" not in report  # the only unmatched quote belongs to a verified claim
    assert "changed since the deck was written" in report and "votes OUTDATED/CRITICAL" in report
    checklist = (root / "deck-a" / "output" / "checklist.md").read_text()
    assert "- [ ] P001 · REBUILD · slide 2-3" in checklist
    assert run("plan", "report", "--manifest", m, "--deck", "deck-b")["ok"]
    r = run("plan", "course-summary", "--manifest", m)
    assert r["ok"] and "significant update" in (root / "course" / "course_summary.md").read_text()

    # annotated PPTX keeps the original notes after the new lines
    r = run("pptx", "annotate", "--manifest", m, "--deck", "deck-a")
    assert r["ok"] and r["slides_annotated"] == 2, r
    from pptx import Presentation
    notes = Presentation(r["out"]).slides[1].notes_slide.notes_text_frame.text
    assert notes.startswith("[SlideGen P001] REBUILD slides 2-3") and notes.rstrip().endswith("Original note 2")


def test_schema_errors_name_the_file_and_field(tmp_path):
    ws = tmp_path / "ws"
    make_pptx(tmp_path / "d.pptx", ["One"])
    run("ingest", tmp_path / "d.pptx", "--out", ws)
    write(ws / "claims" / "extract_a.json", {"deck": "d", "part": "a", "slides": [1, 1], "claims": [
        {"id": "bad-id", "slide": 1, "text": "t", "statement": "s", "kind": ["named"], "from_image": False}]})
    r = run("claims", "merge", ws)
    assert not r["ok"] and "extract_a.json" in r["error"] and "/claims/0/id" in r["error"]


def test_refetch_keeps_earlier_content(tmp_path):
    """A second download of a changed page keeps the first content id resolvable."""
    import json as _json
    from sglib.sources import _record, entries_by_id
    store = tmp_path / "store"
    for cid, text in (("aaa1", "first version"), ("bbb2", "second version")):
        (store / "text").mkdir(parents=True, exist_ok=True)
        (store / "text" / f"{cid}.txt").write_text(text)
        _record(store, {"url": "https://x.example/p", "status": "fetched", "fetched": cid,
                        "text_file": f"text/{cid}.txt", "retrieved": "2026-10-09T00:00:00+00:00"})
    index = _json.loads((store / "index.json").read_text())["sources"]
    by_id = entries_by_id(index)
    assert set(by_id) == {"aaa1", "bbb2"} and by_id["aaa1"]["text_file"] == "text/aaa1.txt"


def test_match_quote_ignores_spaces_around_links():
    """HTML extraction leaves spaces where inline links ended; a quote from the rendered page still matches."""
    text = "The result is a delta-driven dataset . CDC is an approach ( ETL ) to data integration , used widely ."
    assert match_quote("delta-driven dataset. CDC is an approach (ETL) to data integration, used widely.", text)[0]
    assert match_quote("the result is a delta-driven dataset . CDC is an approach", text)[0]


def test_assumptions_and_arithmetic_skip_research(tmp_path):
    """Assumed values are listed but not checked; arithmetic is batched apart and needs no research."""
    ws = tmp_path / "ws"
    claim = lambda n, kind: {"id": f"claim:d:a-{n:03d}", "slide": n, "slide_title": "t", "text": f"text {n}",
                             "statement": f"statement {n}", "kind": [kind], "from_image": False}
    write(ws / "claims.json", {"deck": "d", "claims": [claim(1, "named"), claim(2, "assumption"),
                                                       claim(3, "arithmetic"), claim(4, "general")]})
    r = run("claims", "select", ws, "--batch-size", 10)
    assert r["ok"] and r["batches"] == 1 and r["arithmetic_batches"] == 1 and r["assumptions_not_checked"] == 1, r
    batches = json.loads((ws / "research" / "batches.json").read_text())
    assert [(b["batch"], b["research"], len(b["claims"])) for b in batches["batches"]] == [("b01", True, 2), ("c01", False, 1)]
    verified = {"verdict": "VERIFIED", "severity": None, "fix_kind": None, "outdated": None,
                "issue": None, "suggested_fix": None, "sources_used": []}
    write(ws / "checks" / "b01.json", {"batch": "b01", "verdicts": [
        {"claim_id": "claim:d:a-001", **verified}, {"claim_id": "claim:d:a-004", **verified}]})
    r = run("claims", "summary", ws)
    assert not r["ok"] and "c01" in r["error"], r          # the arithmetic batch still needs its verdicts
    write(ws / "checks" / "c01.json", {"batch": "c01", "verdicts": [{"claim_id": "claim:d:a-003", **verified}]})
    r = run("claims", "summary", ws)
    assert r["ok"] and r["claims"] == 3 and r["assumptions_not_checked"] == 1, r
    check = json.loads((ws / "claim_check.json").read_text())
    assert [a["claim_id"] for a in check["assumptions"]] == ["claim:d:a-002"]
