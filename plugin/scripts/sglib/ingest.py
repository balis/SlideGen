"""Read a deck (PPTX, PDF, or Marp Markdown) into workspace/slides.json.

Every later step works from slides.json, so it records what a reader of the
deck would see: titles, body text with indent levels, tables, code, speaker
notes, hidden slides, and the embedded images converted to PNG.
"""
from __future__ import annotations

import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from . import images

SCHEMA_VERSION = 1
MOSTLY_IMAGE_COVERAGE = 0.5   # share of the slide area covered by pictures
MOSTLY_IMAGE_MAX_TEXT = 200   # characters of text, excluding the title
FULL_SLIDE_IMAGE = 0.9        # a picture this large is the slide, whatever text overlays it
MIN_CODE_CHARS = 40           # a single monospace line shorter than this is inline, not code
OPENING_TEXT_CHARS = 80
MONOSPACE_FONTS = {"consolas", "courier", "courier new", "menlo", "monaco",
                   "source code pro", "fira code", "jetbrains mono", "lucida console"}


def _clean(text: str) -> str:
    return re.sub(r"[ \t]+", " ", text.replace("\x0b", " ").replace("\xa0", " ")).strip()


def _picture(shape) -> tuple[bytes, str]:
    """A picture's bytes and file extension. python-pptx's Image.ext asks Pillow to identify the image, which fails
    for some WMF and EMF files; the extension of the image's package part is used then."""
    image = shape.image
    try:
        ext = image.ext
    except Exception:   # PIL.UnidentifiedImageError and decoder errors
        ext = shape.part.related_part(shape._element.blip_rId).partname.ext
    return image.blob, ext.lower()


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _finish_slide(slide: dict) -> dict:
    """Derive opening text and the mostly-image flag from the collected content."""
    body_text = " ".join(p["text"] for p in slide["body"])
    slide["opening_text"] = body_text[:OPENING_TEXT_CHARS]
    slide["text_chars"] = len(body_text)
    slide["mostly_image"] = (slide["image_coverage"] >= FULL_SLIDE_IMAGE
                             or (slide["image_coverage"] >= MOSTLY_IMAGE_COVERAGE
                                 and slide["text_chars"] < MOSTLY_IMAGE_MAX_TEXT))
    if slide["title"] is None and not slide["mostly_image"] and slide["body"]:
        # Title slides and free-form layouts: the first line is the best available title.
        slide["title"], slide["title_source"] = slide["body"][0]["text"], "first_text"
    return slide


def _new_slide(index: int) -> dict:
    return {"index": index, "hidden": False, "title": None, "title_source": None,
            "body": [], "code_blocks": [], "tables": [], "charts": [], "notes": "",
            "images": [], "image_coverage": 0.0}


class _ImageStore:
    """Saves each distinct embedded image once, as PNG, under workspace/images/."""

    def __init__(self, workspace: Path):
        self.dir = workspace / "images"
        self.by_hash: dict[str, dict] = {}
        self.failed: list[dict] = []

    def add(self, blob: bytes, ext: str, slide_index: int, coverage: float) -> dict:
        digest = hashlib.sha1(blob).hexdigest()[:12]
        if digest not in self.by_hash:
            image, fmt = images.to_png(blob, ext)
            entry = {"file": None, "source_format": fmt, "size": None}
            if image is not None:
                rel = Path("images") / f"s{slide_index:03d}_{digest}.png"
                entry["size"] = list(images.save_png(image, self.dir.parent / rel))
                entry["file"] = str(rel)
            else:
                self.failed.append({"slide": slide_index, "format": fmt})
            self.by_hash[digest] = entry
        return {**self.by_hash[digest], "coverage": round(coverage, 3)}


# ---------------------------------------------------------------- PPTX

def _read_pptx(path: Path, store: _ImageStore) -> list[dict]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

    skip_placeholders = {PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE}
    title_placeholders = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}

    prs = Presentation(str(path))
    area = (prs.slide_width or 1) * (prs.slide_height or 1)
    height = prs.slide_height or 1

    def walk(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(shape.shapes)
            else:
                yield shape

    def placeholder_type(shape):
        return shape.placeholder_format.type if shape.is_placeholder else None

    slides = []
    for index, sl in enumerate(prs.slides, 1):
        slide = _new_slide(index)
        slide["hidden"] = sl._element.get("show") == "0"
        shapes = [s for s in walk(sl.shapes) if placeholder_type(s) not in skip_placeholders]
        shapes.sort(key=lambda s: ((s.top or 0), (s.left or 0)))

        title_shape = next((s for s in shapes if placeholder_type(s) in title_placeholders
                            and s.has_text_frame and _clean(s.text_frame.text)), None)
        if title_shape is not None:
            slide["title"], slide["title_source"] = _clean(title_shape.text_frame.text), "placeholder"
        else:
            # Fall back to the top-most short text in the upper quarter of the slide.
            for s in shapes:
                text = _clean(s.text_frame.text) if s.has_text_frame else ""
                if text and len(text) <= 120 and (s.top or 0) < height / 4:
                    title_shape, slide["title"], slide["title_source"] = s, text, "top_text"
                    break

        coverage = 0.0
        code: list[str] = []

        def flush_code():
            if len(code) > 1 or (code and len(code[0]) >= MIN_CODE_CHARS):
                slide["code_blocks"].append("\n".join(code))
            else:
                slide["body"].extend({"text": line, "level": 0} for line in code)
            code.clear()

        for s in shapes:
            if s is title_shape:
                continue
            blob = None
            try:
                if s.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(s, "image"):
                    blob, ext = _picture(s)
            except (AttributeError, ValueError, KeyError):
                blob = None
            if blob:
                share = min(1.0, ((s.width or 0) * (s.height or 0)) / area)
                coverage += share
                slide["images"].append(store.add(blob, ext, index, share))
            if getattr(s, "has_table", False) and s.has_table:
                slide["tables"].append([[_clean(c.text) for c in row.cells] for row in s.table.rows])
            if getattr(s, "has_chart", False) and s.has_chart:
                chart = s.chart
                try:
                    series = [ser.name for plot in chart.plots for ser in plot.series]
                except (KeyError, ValueError, NotImplementedError):
                    series = []
                slide["charts"].append({
                    "type": str(chart.chart_type),
                    "title": _clean(chart.chart_title.text_frame.text) if chart.has_title else None,
                    "series": series})
            if s.has_text_frame:
                for para in s.text_frame.paragraphs:
                    text = _clean("".join(r.text for r in para.runs) or para.text)
                    if not text:
                        continue
                    fonts = {(r.font.name or "").lower() for r in para.runs if r.text.strip()}
                    if fonts and fonts <= MONOSPACE_FONTS:
                        code.append(text)
                        continue
                    flush_code()
                    slide["body"].append({"text": text, "level": para.level})
                flush_code()  # a code block never spans two shapes
        slide["image_coverage"] = round(min(coverage, 1.0), 3)
        if sl.has_notes_slide:
            slide["notes"] = sl.notes_slide.notes_text_frame.text.strip()
        slides.append(_finish_slide(slide))
    return slides


# ---------------------------------------------------------------- PDF

def _read_pdf(path: Path, store: _ImageStore) -> list[dict]:
    import fitz  # PyMuPDF

    slides = []
    with fitz.open(str(path)) as doc:
        for index, page in enumerate(doc, 1):
            slide = _new_slide(index)
            page_area = page.rect.width * page.rect.height or 1
            lines = []
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    text = _clean("".join(span["text"] for span in line["spans"]))
                    if text:
                        size = max(span["size"] for span in line["spans"])
                        lines.append((line["bbox"][1], size, text))
            if lines:
                top = min(lines, key=lambda l: (-l[1], l[0]))  # largest font, then top-most
                slide["title"], slide["title_source"] = top[2], "largest_text"
                slide["body"] = [{"text": t, "level": 0} for _, _, t in sorted(lines) if t != top[2]]
            coverage = sum(min(1.0, (r.width * r.height) / page_area)
                           for info in page.get_image_info() for r in [fitz.Rect(info["bbox"])])
            slide["image_coverage"] = round(min(coverage, 1.0), 3)
            _finish_slide(slide)
            if slide["mostly_image"]:
                # Render the whole page: it shows the slide exactly as the audience saw it.
                pix = page.get_pixmap(dpi=150)
                slide["images"].append(store.add(pix.tobytes("png"), "png", index, 1.0))
            slides.append(slide)
    return slides


# ---------------------------------------------------------------- Marp

_COMMENT = re.compile(r"<!--(.*?)-->", re.S)
_DIRECTIVE = re.compile(r"^\s*_?[a-zA-Z]+\s*:")


def _read_marp(path: Path, store: _ImageStore) -> list[dict]:
    text = path.read_text(encoding="utf-8")
    if text.startswith("---\n"):
        end = text.find("\n---\n", 4)
        text = text[end + 5:] if end != -1 else text
    # Split on slide separators that are not inside code fences.
    chunks, current, in_fence = [], [], False
    for line in text.splitlines():
        if line.strip().startswith("```"):
            in_fence = not in_fence
        if line.strip() == "---" and not in_fence:
            chunks.append("\n".join(current))
            current = []
        else:
            current.append(line)
    chunks.append("\n".join(current))

    slides = []
    for index, chunk in enumerate((c for c in chunks if c.strip()), 1):
        slide = _new_slide(index)
        notes = [c.strip() for c in _COMMENT.findall(chunk) if not _DIRECTIVE.match(c)]
        slide["notes"] = "\n\n".join(notes)
        body_md = _COMMENT.sub("", chunk)
        for fence in re.findall(r"```[^\n]*\n(.*?)```", body_md, re.S):
            slide["code_blocks"].append(fence.rstrip())
        body_md = re.sub(r"```.*?```", "", body_md, flags=re.S)
        for line in body_md.splitlines():
            if not line.strip():
                continue
            heading = re.match(r"^(#+)\s+(.*)", line)
            if heading and slide["title"] is None:
                slide["title"], slide["title_source"] = _clean(heading.group(2)), "heading"
                continue
            bullet = re.match(r"^(\s*)[-*+]\s+(.*)", line)
            level = len(bullet.group(1)) // 2 if bullet else 0
            slide["body"].append({"text": _clean(bullet.group(2) if bullet else line), "level": level})
        slides.append(_finish_slide(slide))
    return slides


# ---------------------------------------------------------------- command

READERS = {".pptx": _read_pptx, ".pdf": _read_pdf, ".md": _read_marp}


def run(args) -> dict:
    deck = Path(args.deck).resolve()
    workspace = Path(args.out).resolve()
    out_file = workspace / "slides.json"
    reader = READERS.get(deck.suffix.lower())
    if reader is None:
        return {"ok": False, "error": f"unsupported deck format: {deck.suffix}"}
    if not deck.exists():
        return {"ok": False, "error": f"deck not found: {deck}"}

    digest = _sha256(deck)
    if out_file.exists() and not args.force:
        existing = json.loads(out_file.read_text(encoding="utf-8"))
        if existing.get("source_sha256") == digest:
            return {"ok": True, "skipped": True, "out": str(out_file),
                    "reason": "slides.json is up to date for this deck"}

    workspace.mkdir(parents=True, exist_ok=True)
    store = _ImageStore(workspace)
    slides = reader(deck, store)

    if args.pdf:  # Render image-heavy slides from the PDF export as well.
        pdf_store = _ImageStore(workspace)
        pdf_slides = _read_pdf(Path(args.pdf).resolve(), pdf_store)
        visible = [s for s in slides if not s["hidden"]]  # PDF exports omit hidden slides
        for slide, page in zip(visible, pdf_slides):
            if slide["mostly_image"] or page["mostly_image"]:
                import fitz
                with fitz.open(args.pdf) as doc:
                    pix = doc[page["index"] - 1].get_pixmap(dpi=150)
                rendered = pdf_store.add(pix.tobytes("png"), "png", slide["index"], 1.0)
                slide["rendered"] = rendered["file"]

    result = {
        "schema_version": SCHEMA_VERSION,
        "source": str(deck),
        "source_sha256": digest,
        "format": deck.suffix.lower().lstrip("."),
        "ingested_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "slides": slides,
    }
    out_file.write_text(json.dumps(result, ensure_ascii=False, indent=1), encoding="utf-8")
    return {
        "ok": not store.failed,
        "out": str(out_file),
        "slides": len(slides),
        "hidden": sum(s["hidden"] for s in slides),
        "mostly_image": sum(s["mostly_image"] for s in slides),
        "untitled": sum(s["title"] is None for s in slides),
        "images_saved": sum(1 for e in store.by_hash.values() if e["file"]),
        "images_failed": store.failed,
    }
