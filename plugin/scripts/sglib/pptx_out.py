"""PowerPoint output: a copy of the deck with each slide's plan items in its speaker notes."""
from __future__ import annotations

import shutil
from pathlib import Path

from .plan import _Deck
from .util import SgError

MARK = "[SlideGen"


def _note_lines(item: dict, slide: int) -> list[str]:
    span = item["slides"]
    if item["type"] == "INSERT":
        head = f"{MARK} {item['id']}] INSERT after this slide"
    elif item["type"] == "REBUILD":
        head = f"{MARK} {item['id']}] REBUILD slides {span[0]}-{span[-1]} (this is slide {slide})"
    else:
        head = f"{MARK} {item['id']}] {item['type']}"
    lines = [f"{head}: {item['reason']}"]
    if item["current_text"] and item["type"] not in {"REBUILD", "INSERT"}:
        lines.append(f"  Now: {item['current_text']}")
    if item["type"] != "REBUILD" or slide == span[0]:  # print a rebuild's outline once
        lines.append(f"  Proposed: {item['proposed']}")
    return lines


def _prepend(frame, lines: list[str]) -> None:
    """Insert lines as new paragraphs before the existing notes, keeping their formatting."""
    had_text = bool(frame.text.strip())
    body = frame._txBody
    first_p = body.find("{http://schemas.openxmlformats.org/drawingml/2006/main}p")
    new = []
    for line in lines + ([""] if had_text else []):
        para = frame.add_paragraph()   # appended at the end; moved below
        para.text = line
        new.append(para._p)
    for p in new:
        body.remove(p)
    position = list(body).index(first_p) if first_p is not None else len(body)
    if not had_text and first_p is not None:  # an empty placeholder paragraph: replace it
        body.remove(first_p)
    for offset, p in enumerate(new):
        body.insert(position + offset, p)


def annotate(args) -> dict:
    deck = _Deck(args.manifest, args.deck)
    if deck.plan is None:
        raise SgError("no plan.json; run the change-planner and `sg plan validate` first")
    src = Path(deck.cfg["_path"])
    if src.suffix.lower() != ".pptx":
        raise SgError(f"annotate needs a PPTX deck, got {src.name}")
    approved = None
    if args.approved:
        approved = {line.split()[2] for line in Path(args.approved).read_text(encoding="utf-8").splitlines()
                    if line.startswith("- [x] ")}

    from pptx import Presentation
    deck.out.mkdir(parents=True, exist_ok=True)
    dst = deck.out / "annotated.pptx"
    shutil.copyfile(src, dst)
    prs = Presentation(str(dst))
    per_slide: dict[int, list[str]] = {}
    for item in sorted(deck.plan["items"], key=lambda i: i["id"]):
        if approved is not None and item["id"] not in approved:
            continue
        targets = item["slides"] if item["type"] in {"REBUILD", "MERGE", "COMPRESS"} else item["slides"][:1]
        for s in targets:
            per_slide.setdefault(s, []).extend(_note_lines(item, s))
    for index, slide in enumerate(prs.slides, 1):
        lines = per_slide.get(index)
        if lines:
            _prepend(slide.notes_slide.notes_text_frame, lines)
    prs.save(str(dst))
    return {"ok": True, "out": str(dst), "slides_annotated": len(per_slide),
            "items": sum(1 for i in deck.plan["items"] if approved is None or i["id"] in approved)}
