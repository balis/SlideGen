"""Build a lecture deck as native PowerPoint slides on an existing deck's master.

A deck specification (schema `deck_spec`) lists the slides in order. A slide either copies a slide of the base deck
(with text edits) or is new: a layout of the base deck's master, a title, bullets, speaker notes, and diagram shapes
drawn as editable PowerPoint shapes. Slides of the base deck that the specification does not copy are removed, so
the output keeps the base deck's master, theme, fonts and footers.

  sg pptx build --spec deck.json          writes the .pptx named in the specification
  sg pptx render --pptx deck.pptx --out DIR   renders slides to PNG with LibreOffice (outside the sandbox)
"""
from __future__ import annotations

import copy
import os
import re
import shutil
import subprocess
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .util import SgError, load_json

PALETTE = {"blue": "4472C4", "orange": "ED7D31", "grey": "A5A5A5", "gold": "FFC000", "lightblue": "5B9BD5",
           "green": "70AD47", "dark": "44546A", "light": "E7E6E6", "white": "FFFFFF", "black": "000000",
           "red": "C00000", "paleblue": "DAE3F3", "paleorange": "FBE5D6", "palegreen": "E2F0D9", "palegrey": "EDEDED",
           "palegold": "FFF2CC"}
SHAPES = {"rect": MSO_SHAPE.RECTANGLE, "rounded": MSO_SHAPE.ROUNDED_RECTANGLE, "ellipse": MSO_SHAPE.OVAL,
          "can": MSO_SHAPE.CAN, "chevron": MSO_SHAPE.CHEVRON, "pentagon": MSO_SHAPE.PENTAGON,
          "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE}
ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
MARKUP = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*")   # **bold**, *italic*


def _rgb(name: str | None, palette: dict) -> RGBColor | None:
    if not name:
        return None
    hexval = palette.get(name, name).lstrip("#")
    if not re.fullmatch(r"[0-9A-Fa-f]{6}", hexval):
        raise SgError(f"unknown colour {name}")
    return RGBColor.from_string(hexval.upper())


def _runs(paragraph, text: str, size=None, color=None, bold=False, italic=False) -> None:
    """Fill a paragraph with text, where **...** marks bold and *...* italic runs."""
    pos = 0
    parts = []
    for m in MARKUP.finditer(text):
        if m.start() > pos:
            parts.append((text[pos:m.start()], bold, italic))
        parts.append((m.group(1), True, italic) if m.group(1) is not None else (m.group(2), bold, True))
        pos = m.end()
    if pos < len(text):
        parts.append((text[pos:], bold, italic))
    for chunk, b, it in parts or [("", bold, italic)]:
        run = paragraph.add_run()
        run.text = chunk
        run.font.bold = b or None
        run.font.italic = it or None
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color


def _bullet(p, level: int) -> None:
    """A bullet with a hanging indent, for text boxes (placeholders take bullets from the master)."""
    ppr = p._p.get_or_add_pPr()
    indent = 228600   # 0.25 in
    ppr.set("marL", str(indent * (level + 1)))
    ppr.set("indent", str(-indent))
    spc = etree.SubElement(ppr, qn("a:spcBef"))
    etree.SubElement(spc, qn("a:spcPts")).set("val", "400")
    etree.SubElement(ppr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(ppr, qn("a:buChar")).set("char", "•" if level == 0 else "–")


def _fill_frame(frame, lines: list, palette: dict, size=None, color=None, align=None, bold=False, italic=False,
                bullets=False):
    """Write lines (strings or {"text", "level"}) into a text frame, replacing its content."""
    frame.clear()
    for i, line in enumerate(lines):
        item = line if isinstance(line, dict) else {"text": line}
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        if bullets and item.get("bullet", True):
            _bullet(p, item.get("level", 0))
        else:
            p.level = item.get("level", 0)
        if align:
            p.alignment = ALIGN[align]
        _runs(p, item["text"], item.get("size", size), _rgb(item.get("color"), palette) or color,
              item.get("bold", bold), item.get("italic", italic))


def _slide_number(slide) -> None:
    """Give a new slide the layout's slide-number placeholder with a live slide-number field."""
    for ph in slide.slide_layout.placeholders:
        if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            sp = slide.shapes.clone_placeholder(ph)
            shape = slide.shapes[-1]
            body = shape.text_frame._txBody
            for p in body.findall(qn("a:p")):
                body.remove(p)
            p = etree.SubElement(body, qn("a:p"))
            fld = etree.SubElement(p, qn("a:fld"), id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum")
            etree.SubElement(fld, qn("a:t")).text = "‹#›"
            return sp


def _arrowheads(connector, head: str) -> None:
    ln = connector.line._get_or_add_ln()
    for tag, wanted in (("a:headEnd", head in ("start", "both")), ("a:tailEnd", head in ("end", "both"))):
        if wanted:
            el = etree.SubElement(ln, qn(tag))
            el.set("type", "triangle")
            el.set("w", "med")
            el.set("len", "med")


def _draw(slide, shp: dict, palette: dict) -> None:
    t = shp["type"]
    if t in ("arrow", "line"):
        (x1, y1), (x2, y2) = shp["from"], shp["to"]
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = _rgb(shp.get("color", "dark"), palette)
        c.line.width = Pt(shp.get("width", 1.5))
        if shp.get("dashed"):
            c.line.dash_style = 4   # MSO_LINE_DASH_STYLE.DASH
        _arrowheads(c, shp.get("head", "end" if t == "arrow" else "none"))
        if shp.get("label"):
            lx, ly = (x1 + x2) / 2, (y1 + y2) / 2
            _draw(slide, {"type": "text", "x": lx - 0.9, "y": ly - 0.3, "w": 1.8, "h": 0.25, "text": shp["label"],
                          "size": shp.get("label_size", 10), "align": "center", "color": shp.get("color", "dark")},
                  palette)
        return
    x, y, w, h = (Inches(shp[k]) for k in ("x", "y", "w", "h"))
    if t == "text":
        box = slide.shapes.add_textbox(x, y, w, h)
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(0.03)
        lines = shp["text"] if isinstance(shp["text"], list) else shp["text"].split("\n")
        _fill_frame(frame, lines, palette, shp.get("size", 14), _rgb(shp.get("color"), palette), shp.get("align"),
                    shp.get("bold", False), shp.get("italic", False), shp.get("bullets", False))
        frame.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                                 "bottom": MSO_ANCHOR.BOTTOM}[shp.get("valign", "top")]
        return
    if t == "table":
        rows = shp["rows"]
        gt = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, h).table
        if shp.get("col_widths"):
            for i, cw in enumerate(shp["col_widths"]):
                gt.columns[i].width = Inches(cw)
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = gt.cell(r, c)
                _fill_frame(cell.text_frame, str(val).split("\n"), palette, shp.get("size", 12),
                            bold=(r == 0 and shp.get("header", True)))
                cell.margin_top = cell.margin_bottom = Inches(0.03)
        return
    if t == "box":
        box = slide.shapes.add_shape(SHAPES[shp.get("shape", "rect")], x, y, w, h)
        if shp.get("fill") == "none":
            box.fill.background()
        else:
            box.fill.solid()
            box.fill.fore_color.rgb = _rgb(shp.get("fill", "paleblue"), palette)
        if shp.get("line") == "none":
            box.line.fill.background()
        else:
            box.line.color.rgb = _rgb(shp.get("line", "dark"), palette)
            box.line.width = Pt(shp.get("line_width", 1))
            if shp.get("dashed"):
                box.line.dash_style = 4
        box.shadow.inherit = False
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(0.05)
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        lines = shp.get("text", "")
        lines = lines if isinstance(lines, list) else lines.split("\n")
        _fill_frame(frame, lines, palette, shp.get("size", 12), _rgb(shp.get("text_color", "black"), palette),
                    shp.get("align", "center"), shp.get("bold", False))
        return
    raise SgError(f"unknown shape type {t}")


def _replace(slide, find: str, replace: str) -> bool:
    """Replace text on a slide, changing only the runs the found text spans: the replacement takes the formatting of
    the run where the found text starts, and runs before and after it keep theirs."""
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for p in shape.text_frame.paragraphs:
            runs = p.runs
            full = "".join(r.text for r in runs)
            at = full.find(find)
            if at < 0:
                continue
            stop, pos = at + len(find), 0
            spans = []
            for r in runs:
                spans.append((r, pos, pos + len(r.text)))
                pos += len(r.text)
            hit = [(r, a, b) for r, a, b in spans if a < stop and b > at]
            first, fa, _ = hit[0]
            last, la, lb = hit[-1]
            tail = last.text[stop - la:]
            first.text = first.text[:at - fa] + replace + (tail if first is last else "")
            if first is not last:
                last.text = tail
                for r, _, _ in hit[1:-1]:
                    r._r.getparent().remove(r._r)
                if not tail:
                    last._r.getparent().remove(last._r)
            return True
    return False


def _insert_after(slide, after: str, lines: list) -> bool:
    """Insert lines after the paragraph that contains `after`, in that paragraph's style and level (unless a line
    sets its own level)."""
    from pptx.text.text import _Paragraph
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            if after in "".join(r.text for r in para.runs):
                prev = para
                for line in lines:
                    item = line if isinstance(line, dict) else {"text": line}
                    new = copy.deepcopy(prev._p)
                    for r in new.findall(qn("a:r")) + new.findall(qn("a:fld")) + new.findall(qn("a:br")):
                        new.remove(r)
                    prev._p.addnext(new)
                    p = _Paragraph(new, prev._parent)
                    if "level" in item:
                        p.level = item["level"]
                    _runs(p, item["text"], item.get("size"))
                    prev = p
                return True
    return False


def _append_bullets(slide, bullets: list, palette: dict) -> None:
    body = next((s for s in slide.placeholders if s.placeholder_format.idx == 1), None)
    if body is None:
        raise SgError("append_bullets needs a slide with a body placeholder")
    last = body.text_frame.paragraphs[-1]
    for b in bullets:
        item = b if isinstance(b, dict) else {"text": b}
        new = copy.deepcopy(last._p)
        for r in new.findall(qn("a:r")) + new.findall(qn("a:fld")):
            new.remove(r)
        last._p.addnext(new)
        from pptx.text.text import _Paragraph
        p = _Paragraph(new, last._parent)
        p.level = item.get("level", 0)
        _runs(p, item["text"], item.get("size"))
        last = p


def build(args) -> dict:
    spec_path = Path(args.spec).resolve()
    spec = load_json(spec_path, "deck_spec")
    base = (spec_path.parent / spec["base"]).resolve()
    out = (spec_path.parent / spec["output"]).resolve()
    palette = {**PALETTE, **spec.get("palette", {})}
    prs = Presentation(str(base))
    originals = list(prs.slides)
    sld_ids = prs.slides._sldIdLst
    original_ids = [s.slide_id for s in originals]
    id_of = dict(zip(original_ids, list(sld_ids)))
    order, problems, kept = [], [], set()
    for n, item in enumerate(spec["slides"], 1):
        if "copy" in item:
            src = item["copy"]
            if not 1 <= src <= len(originals):
                raise SgError(f"slide {n}: copy {src} is outside the base deck (1..{len(originals)})")
            if src in kept:
                raise SgError(f"slide {n}: base slide {src} is copied twice")
            kept.add(src)
            slide = originals[src - 1]
            slide._element.attrib.pop("show", None)   # a copied slide is shown even if hidden in the base deck
            for e in item.get("edits", []):
                if not _replace(slide, e["find"], e["replace"]):
                    problems.append(f"slide {n} (base {src}): text not found: {e['find'][:60]}")
            if item.get("title"):
                slide.shapes.title.text_frame.text = item["title"]
            for ins in item.get("insert", []):
                if not _insert_after(slide, ins["after"], ins["lines"]):
                    problems.append(f"slide {n} (base {src}): text to insert after not found: {ins['after'][:60]}")
            if item.get("append_bullets"):
                _append_bullets(slide, item["append_bullets"], palette)
            for shp in item.get("shapes", []):
                _draw(slide, shp, palette)
        else:
            layouts = {l.name: l for l in prs.slide_layouts}
            if item["layout"] not in layouts:
                raise SgError(f"slide {n}: unknown layout {item['layout']}; known: {sorted(layouts)}")
            slide = prs.slides.add_slide(layouts[item["layout"]])
            if item.get("title") is not None and slide.shapes.title is not None:
                slide.shapes.title.text_frame.text = item["title"]
            body = next((s for s in slide.placeholders if s.placeholder_format.idx == 1), None)
            if item.get("bullets"):
                if body is None:
                    raise SgError(f"slide {n}: layout {item['layout']} has no body placeholder for bullets")
                _fill_frame(body.text_frame, item["bullets"], palette, item.get("bullet_size"))
                if item.get("body_box"):
                    bx, by, bw, bh = item["body_box"]
                    body.left, body.top, body.width, body.height = Inches(bx), Inches(by), Inches(bw), Inches(bh)
            elif body is not None:
                body._element.getparent().remove(body._element)   # no empty "click to add text" boxes
            for idx, lines in (item.get("columns") or {}).items():
                ph = next((s for s in slide.placeholders if s.placeholder_format.idx == int(idx)), None)
                if ph is None:
                    raise SgError(f"slide {n}: layout {item['layout']} has no placeholder {idx}")
                _fill_frame(ph.text_frame, lines, palette, item.get("bullet_size"))
            for ph in list(slide.placeholders):   # remove placeholders the slide leaves empty
                if ph.placeholder_format.idx not in (0, 1) and not ph.text_frame.text.strip():
                    ph._element.getparent().remove(ph._element)
            _slide_number(slide)
            for shp in item.get("shapes", []):
                _draw(slide, shp, palette)
        if item.get("notes") is not None:
            frame = slide.notes_slide.notes_text_frame
            text = item["notes"] if item.get("notes_mode", "replace") == "replace" or not frame.text.strip() \
                else frame.text + "\n\n" + item["notes"]
            frame.text = text
        order.append(slide.slide_id)
    # reorder: the specification's order; drop base slides that were not copied
    all_ids = {s.slide_id: el for s, el in zip(prs.slides, list(sld_ids))}
    for el in list(sld_ids):
        sld_ids.remove(el)
    for sid in order:
        sld_ids.append(all_ids[sid])
    for sid in original_ids:
        if sid not in order:
            prs.part.drop_rel(id_of[sid].get(qn("r:id")))
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return {"ok": not problems, "slides": len(order), "out": str(out), "problems": problems}


def render(args) -> dict:
    """Render a deck to one PNG per shown slide with LibreOffice, using Windows fonts when they are available."""
    import fitz
    pptx = Path(args.pptx).resolve()
    out = Path(args.out).resolve()
    out.mkdir(parents=True, exist_ok=True)
    env = dict(os.environ)
    conf = out / "fonts.conf"
    win_fonts = Path("/mnt/c/Windows/Fonts")
    conf.write_text("<?xml version='1.0'?><fontconfig><include ignore_missing='yes'>/etc/fonts/fonts.conf</include>"
                    + (f"<dir>{win_fonts}</dir>" if win_fonts.exists() else "")
                    + f"<cachedir>{out / 'fc-cache'}</cachedir></fontconfig>", encoding="utf-8")
    env["FONTCONFIG_FILE"] = str(conf)
    shutil.copyfile(pptx, out / "deck.pptx")
    cmd = ["soffice", f"-env:UserInstallation=file://{out / 'lo-profile'}", "--headless", "--norestore",
           "--convert-to", "pdf", "--outdir", str(out), str(out / "deck.pptx")]
    run = subprocess.run(cmd, env=env, capture_output=True, text=True, timeout=600)
    pdf = out / "deck.pdf"
    if not pdf.exists():
        raise SgError(f"LibreOffice did not write a PDF (run outside the sandbox?): {run.stderr[-300:]}")
    doc = fitz.open(str(pdf))
    pages = []
    for i, page in enumerate(doc, 1):
        path = out / f"slide-{i:02d}.png"
        page.get_pixmap(dpi=args.dpi).save(str(path))
        pages.append(str(path))
    return {"ok": True, "slides": len(pages), "out": str(out), "windows_fonts": win_fonts.exists()}
