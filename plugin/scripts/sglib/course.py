"""Course planning: a reference curriculum, a gap analysis, and a lecture plan that fits the course.

Layout, relative to the directory of the plan manifest (plan.yaml):
  plan/references/<kind>.json   reference-scout output, one file per kind of source
  plan/references_check.json    accepted and rejected references (sg course refs)
  plan/taxonomy.json            taxonomist output
  plan/reference_curriculum.*   topic coverage across references (sg course taxonomy)
  plan/mapping.json             curriculum-mapper output: the current lectures as topics
  plan/gap_analysis.*           gaps, over-weighted topics, overlaps (sg course gaps)
  plan/trend_map.json           trend-mapper output: the field's trend topics linked to course topics
  plan/trends.json|trend_gaps.md   attention, maturity and use per course topic (sg course trends)
  plan/proposal.json            course-architect output
  plan/review.json              course-reviewer output (optional)
  plan/curriculum_proposal.md, plan/decisions.md   rendered by sg course report
  sources/                      downloaded sources, shared with the other skills
"""
from __future__ import annotations

from collections import defaultdict
from pathlib import Path

import yaml

from .sources import entries_by_id, match_quote
from .util import SgError, load_json, today, write_json

KINDS = ("syllabus", "textbook", "paper", "professional")
CORE_SHARE = 0.5          # covered by at least this share of independent references
COMMON_SHARE = 0.25
MIN_TOPIC_QUOTE = 6       # topic quotes are short: lecture or chapter titles
MIN_MATCHED_SHARE = 0.7   # a reference counts only if most of its topic quotes are found
MAX_DECISIONS = 12        # more than this overwhelms the author; group related changes
MAX_TOPICS_PER_LABEL = 3  # a combined syllabus line may name a few distinct topics
TIERS = ("core", "common", "niche")
SUGGESTED_RANGE = (0.88, 1.05)  # around the median size of the author's decks
MAX_MINUTES_OFF = 5       # a lecture's topic hours and its material must agree within this many minutes
SLOT_SLACK = 1            # minutes of rounding allowed when a lecture's material is checked against its slot
NAMED_STEADY = 3          # steady trend topics named per course topic in trend_gaps.md


def _plan(path: str) -> dict:
    p = Path(path).resolve()
    cfg = yaml.safe_load(p.read_text(encoding="utf-8"))
    cfg["_root"], cfg["_ws"], cfg["_store"] = p.parent, p.parent / "plan", p.parent / "sources"
    cfg.setdefault("references", {})
    return cfg


def _store_text(store: Path, content_id: str) -> str | None:
    index_file = store / "index.json"
    if not index_file.exists():
        return None
    entry = entries_by_id(load_json(index_file)["sources"]).get(content_id)
    return (store / entry["text_file"]).read_text(encoding="utf-8") if entry else None


def visible_slides(path: Path) -> int:
    from pptx import Presentation
    return sum(1 for sl in Presentation(str(path)).slides if sl._element.get("show") != "0")


def slide_range(cfg: dict) -> tuple[int, int] | None:
    """The author's expected slide count per lecture (plan.yaml lectures.slides)."""
    rng = cfg["lectures"].get("slides")
    return (int(rng[0]), int(rng[1])) if rng else None


def suggested_slides(cfg: dict) -> list[int] | None:
    """A slide range from the author's own decks: around the median visible-slide count."""
    counts = [visible_slides(cfg["_root"] / d["path"]) for d in cfg.get("decks", [])
              if str(d["path"]).lower().endswith(".pptx") and (cfg["_root"] / d["path"]).exists()]
    if not counts:
        return None
    counts.sort()
    median = (counts[(len(counts) - 1) // 2] + counts[len(counts) // 2]) / 2
    return [round(median * SUGGESTED_RANGE[0]), round(median * SUGGESTED_RANGE[1])]


def slides_per_minute(cfg: dict) -> float | None:
    rng = slide_range(cfg)
    return (rng[0] + rng[1]) / 2 / (cfg["lectures"]["hours"] * 60) if rng else None


def slide_estimate(cfg: dict, lec: dict) -> int | None:
    """Reused slides plus the new material's minutes at the author's pace."""
    pace = slides_per_minute(cfg)
    if pace is None:
        return None
    reused = sum(len(r["slides"]) for r in lec.get("reuse", []))
    return round(reused + sum(m["minutes"] for m in lec.get("new_material", [])) * pace)


def material_minutes(cfg: dict, lec: dict) -> float | None:
    """The time a lecture's material takes: reused slides at the author's pace plus the new material's minutes."""
    pace = slides_per_minute(cfg)
    if pace is None:
        return None
    reused = sum(len(r["slides"]) for r in lec.get("reuse", []))
    return reused / pace + sum(m["minutes"] for m in lec.get("new_material", []))


# ---------------------------------------------------------------- references

def refs(args) -> dict:
    """Check every reference: quotes found in the downloaded page, recency, quotas, vendor share."""
    cfg = _plan(args.plan)
    ws, store = cfg["_ws"], cfg["_store"]
    quotas = {k: cfg["references"].get(k, 0) for k in KINDS}
    recent_years = cfg["references"].get("recent_years", 3)
    max_vendor = cfg["references"].get("max_vendor_share", 0.3)
    year_now = int(today()[:4])

    files = sorted((ws / "references").glob("*.json"))
    if not files:
        raise SgError(f"no reference files in {ws / 'references'}")
    seen_ids, seen_orgs = {}, {}
    accepted, rejected = [], []
    for path in files:
        doc = load_json(path, "references")
        for ref in doc["references"]:
            if ref["id"] in seen_ids:
                raise SgError(f"{path.name}: duplicate reference id {ref['id']} (also in {seen_ids[ref['id']]})")
            seen_ids[ref["id"]] = path.name
            org_key = (ref["kind"], ref["org"].strip().lower())
            reasons = []
            # One course or programme per institution; textbooks and papers share publishers and venues.
            if ref["kind"] in ("syllabus", "professional") and org_key in seen_orgs:
                reasons.append(f"second {ref['kind']} from {ref['org']} (first: {seen_orgs[org_key]})")
            if ref["kind"] != "textbook" and ref["year"] < year_now - recent_years:
                reasons.append(f"older than {recent_years} years ({ref['year']})")
            text = _store_text(store, ref["fetched"]) if ref["fetched"] else None
            if text is None:
                reasons.append("page not downloaded with sg fetch")
                matched = 0
            else:
                matched = 0
                for t in ref["topics"]:
                    t["matched"] = match_quote(t["quote"], text, min_chars=MIN_TOPIC_QUOTE)[0]
                    matched += t["matched"]
            share = matched / max(len(ref["topics"]), 1)
            if text is not None and share < MIN_MATCHED_SHARE:
                reasons.append(f"only {matched} of {len(ref['topics'])} topic quotes found in the page")
            if reasons:
                rejected.append({"id": ref["id"], "kind": ref["kind"], "reasons": reasons})
            else:
                seen_orgs[org_key] = ref["id"]
                accepted.append({**ref, "file": path.name, "matched_share": round(share, 2)})

    counts = {k: sum(r["kind"] == k for r in accepted) for k in KINDS}
    vendor_share = sum(r["vendor"] for r in accepted) / max(len(accepted), 1)
    short = {k: f"{counts[k]}/{quotas[k]}" for k in KINDS if counts[k] < quotas[k]}
    problems = []
    if short:
        problems.append(f"too few accepted references: {short}")
    if vendor_share > max_vendor:
        problems.append(f"vendor share {vendor_share:.0%} exceeds {max_vendor:.0%}")
    result = {"generated_by": "sg course refs", "ok": not problems, "problems": problems,
              "counts": counts, "quotas": quotas, "vendor_share": round(vendor_share, 2),
              "accepted": accepted, "rejected": rejected}
    write_json(ws / "references_check.json", result)
    return {"ok": not problems, "out": str(ws / "references_check.json"), "counts": counts,
            "vendor_share": round(vendor_share, 2), "rejected": rejected, "problems": problems}


def _accepted(ws: Path) -> list[dict]:
    check = load_json(ws / "references_check.json")
    if not check["ok"]:
        raise SgError("references_check.json is not ok; fix the references first: " + "; ".join(check["problems"]))
    return check["accepted"]


# ---------------------------------------------------------------- taxonomy and coverage

def _check_acyclic(topics: dict[str, dict]) -> list[str]:
    state, cycles = {}, []

    def visit(tid, path):
        if state.get(tid) == "done":
            return
        if state.get(tid) == "active":
            cycles.append(" -> ".join(path + [tid]))
            return
        state[tid] = "active"
        for dep in topics[tid]["depends_on"]:
            if dep in topics:
                visit(dep, path + [tid])
        state[tid] = "done"

    for tid in topics:
        visit(tid, [])
    return cycles


def _scope(ref: dict) -> str:
    return ref.get("scope") or ("specialist" if ref["kind"] == "paper" else "curriculum")


def _coverage(refs_: list[dict], taxonomy: dict) -> list[dict]:
    """Tiers come from independent curriculum-scope references (whole courses, programmes, general
    textbooks). Specialist references (papers, single-technology books) are too narrow to count in
    coverage; recent ones signal emerging topics instead."""
    topics = {t["id"]: t for t in taxonomy["topics"]}
    by_ref_topic: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    labels: dict[tuple, list[str]] = defaultdict(list)
    for m in taxonomy["mapping"]:
        labels[(m["ref"], m["label"])].append(m["topic"])
    for r in refs_:
        total = sum(t["weight"] for t in r["topics"]) or 1
        for t in r["topics"]:
            targets = labels[(r["id"], t["label"])]
            for topic in targets:   # a combined line ("Kafka and stream processing") counts for each topic
                by_ref_topic[r["id"]][topic] += t["weight"] / total / len(targets)
    year_now = int(today()[:4])
    indep = [r for r in refs_ if not r["vendor"]]
    curriculum = [r for r in indep if _scope(r) == "curriculum"]
    specialist = [r for r in indep if _scope(r) == "specialist" and r["year"] >= year_now - 3]
    professional = [r for r in indep if r["kind"] == "professional" and r["year"] >= year_now - 3]
    n = max(len(curriculum), 1)
    bachelor = [r for r in curriculum if r["level"] == "bachelor"]
    recent = [r for r in curriculum if r["year"] >= year_now - 1]

    rows = []
    for tid, t in topics.items():
        has = lambda r: tid in by_ref_topic[r["id"]]
        with_ = [r for r in curriculum if has(r)]
        cov = len(with_) / n
        n_spec = sum(map(has, specialist))
        n_prof = sum(map(has, professional))
        older_with = [r for r in with_ if r["year"] < year_now - 1]
        if cov < CORE_SHARE and (n_spec >= 2 or (n_spec >= 1 and n_prof >= 1)):
            trend = "emerging"   # a topic most curricula already teach is mainstream, not emerging
        elif with_ and not any(has(r) for r in recent) and len(older_with) >= 2:
            trend = "fading"
        else:
            trend = "steady"
        emphasis = (sum(by_ref_topic[r["id"]][tid] for r in with_) / len(with_)) if with_ else 0.0
        rows.append({
            "topic": tid, "title": t["title"], "area": t["area"],
            "tier": "core" if cov >= CORE_SHARE else "common" if cov >= COMMON_SHARE else "niche",
            "trend": trend, "coverage": round(cov, 3), "references": len(with_), "independent_total": n,
            "bachelor_coverage": (round(sum(map(has, bachelor)) / len(bachelor), 3) if bachelor else None),
            "specialist_references": n_spec, "professional_references": n_prof,
            "emphasis": round(emphasis, 3),
            "vendor_references": sum(1 for r in refs_ if r["vendor"] and has(r)),
            "sources": [r["id"] for r in refs_ if has(r)],
            "depends_on": t["depends_on"]})
    rows.sort(key=lambda r: (TIERS.index(r["tier"]), -r["coverage"], r["topic"]))
    return rows


def taxonomy(args) -> dict:
    """Validate taxonomy.json against the accepted references and compute topic coverage."""
    cfg = _plan(args.plan)
    ws = cfg["_ws"]
    refs_ = _accepted(ws)
    tax = load_json(ws / "taxonomy.json", "taxonomy")
    topics = {t["id"]: t for t in tax["topics"]}
    areas = {a["id"] for a in tax["areas"]}
    problems = []
    if len(topics) != len(tax["topics"]):
        problems.append("duplicate topic ids")
    problems += [f"topic {t['id']} has unknown area {t['area']}" for t in tax["topics"] if t["area"] not in areas]
    problems += [f"topic {t['id']} depends on unknown topic {d}" for t in tax["topics"]
                 for d in t["depends_on"] if d not in topics]
    expected = {(r["id"], t["label"]) for r in refs_ for t in r["topics"]}
    mapped: dict[tuple, list] = defaultdict(list)
    for m in tax["mapping"]:
        mapped[(m["ref"], m["label"])].append(m["topic"])
        if m["topic"] not in topics:
            problems.append(f"label '{m['label']}' of {m['ref']} maps to unknown topic {m['topic']}")
    missing = sorted(expected - set(mapped))
    if missing:
        problems.append(f"{len(missing)} reference topics are not mapped, e.g. {missing[:5]}")
    repeated = [k for k, v in mapped.items() if len(v) != len(set(v))]
    if repeated:
        problems.append(f"{len(repeated)} labels map to the same topic twice, e.g. {repeated[:3]}")
    too_many = [k for k, v in mapped.items() if len(set(v)) > MAX_TOPICS_PER_LABEL]
    if too_many:
        problems.append(f"{len(too_many)} labels map to more than {MAX_TOPICS_PER_LABEL} topics, e.g. {too_many[:3]}")
    unused = sorted(set(topics) - {m["topic"] for m in tax["mapping"]})
    if unused:
        problems.append(f"topics backed by no reference: {unused}")
    problems += [f"dependency cycle: {c}" for c in _check_acyclic(topics)] if not problems else []
    if problems:
        raise SgError("taxonomy.json: " + "; ".join(problems[:10]))

    rows = _coverage(refs_, tax)
    write_json(ws / "reference_curriculum.json", {"generated_by": "sg course taxonomy", "areas": tax["areas"],
                                                  "references": [{**{k: r[k] for k in ("id", "kind", "title", "url", "org", "year", "level", "vendor")}, "scope": _scope(r)} for r in refs_],
                                                  "topics": rows})
    _render_reference(cfg, tax, refs_, rows)
    tiers = {t: sum(r["tier"] == t for r in rows) for t in TIERS}
    return {"ok": True, "topics": len(rows), "tiers": tiers,
            "emerging": [r["topic"] for r in rows if r["trend"] == "emerging"],
            "fading": [r["topic"] for r in rows if r["trend"] == "fading"],
            "out": str(ws / "reference_curriculum.md")}


def _render_reference(cfg: dict, tax: dict, refs_: list[dict], rows: list[dict]) -> None:
    area_title = {a["id"]: a["title"] for a in tax["areas"]}
    n = rows[0]["independent_total"] if rows else 0
    lines = [f"# Reference curriculum: {cfg.get('name_en') or cfg.get('name')}", "",
             f"Generated {today()} by `sg course taxonomy` from {len(refs_)} references. Coverage counts the "
             f"{n} independent curriculum-scope references (courses, programmes, general textbooks).", "",
             "**Tiers:** core = covered by at least half of them; common = at least a quarter; niche = fewer. "
             "**Trend:** emerging = discussed by recent specialist sources (papers, single-technology books), "
             "or by one of them and an independent professional programme while curricula rarely teach it yet; "
             "fading = taught only by curriculum references older than two years.", "",
             "| Topic | Area | Tier | Coverage | Bachelor | Specialist | Trend |", "|---|---|---|---|---|---|---|"]
    for r in rows:
        bach = f"{r['bachelor_coverage']:.0%}" if r["bachelor_coverage"] is not None else "–"
        lines.append(f"| {r['topic']} {r['title']} | {area_title.get(r['area'], r['area'])} | {r['tier']} | "
                     f"{r['references']}/{r['independent_total']} | {bach} | {r['specialist_references']} | {r['trend']} |")
    lines += ["", "## References", "", "| Id | Kind | Scope | Year | Level | Organisation | Vendor |",
              "|---|---|---|---|---|---|---|"]
    for r in sorted(refs_, key=lambda r: (KINDS.index(r["kind"]), r["id"])):
        lines.append(f"| [{r['id']}]({r['url']}) | {r['kind']} | {_scope(r)} | {r['year']} | {r['level']} | "
                     f"{r['org']} | {'yes' if r['vendor'] else 'no'} |")
    (cfg["_ws"] / "reference_curriculum.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


# ---------------------------------------------------------------- gaps

def gaps(args) -> dict:
    """Compare the current curriculum (mapping.json) with the reference curriculum."""
    cfg = _plan(args.plan)
    ws = cfg["_ws"]
    ref_doc = load_json(ws / "reference_curriculum.json")
    rows = {r["topic"]: r for r in ref_doc["topics"]}
    mapping = load_json(ws / "mapping.json", "mapping")
    problems = [f"lecture {lec['n']}: unknown topic {t['topic']}" for lec in mapping["lectures"]
                for t in lec["topics"] if t["topic"] is not None and t["topic"] not in rows]
    if problems:
        raise SgError("mapping.json: " + "; ".join(problems[:10]))

    taught: dict[str, list[int]] = defaultdict(list)
    hours: dict[str, float] = defaultdict(float)
    current_only = []
    for lec in mapping["lectures"]:
        for t in lec["topics"]:
            if t["topic"] is None:
                current_only.append({"lecture": lec["n"], "label": t["label"], "hours": t["hours"]})
            else:
                taught[t["topic"]].append(lec["n"])
                hours[t["topic"]] += t["hours"]
    total = sum(t["hours"] for lec in mapping["lectures"] for t in lec["topics"]) or 1
    missing = [r for r in ref_doc["topics"] if r["tier"] in ("core", "common") and r["topic"] not in taught]
    emerging_missing = [r for r in ref_doc["topics"] if r["trend"] == "emerging" and r["topic"] not in taught
                        and r["tier"] == "niche"]
    overweight = []
    for tid, h in hours.items():
        r = rows[tid]
        share = h / total
        if r["trend"] == "fading" or (r["tier"] == "niche" and share >= 2 * max(r["emphasis"], 0.02)):
            overweight.append({"topic": tid, "title": r["title"], "hours": round(h, 2), "share": round(share, 3),
                               "reference_emphasis": r["emphasis"], "tier": r["tier"], "trend": r["trend"]})
    overlaps = [{"topic": tid, "title": rows[tid]["title"], "lectures": sorted(set(ls))}
                for tid, ls in taught.items() if len(set(ls)) > 1]
    result = {"generated_by": "sg course gaps", "missing": missing, "emerging_missing": emerging_missing,
              "overweight": overweight, "current_only": current_only, "overlaps": overlaps,
              "taught": {k: sorted(set(v)) for k, v in taught.items()}}
    write_json(ws / "gap_analysis.json", result)

    md = [f"# Gap analysis: {cfg.get('name_en') or cfg.get('name')}", "",
          f"Generated {today()} by `sg course gaps`: the current curriculum compared with the reference curriculum.", "",
          f"## Core and common topics the current curriculum does not teach ({len(missing)})", ""]
    md += [f"- **{r['topic']} {r['title']}** ({r['tier']}, {r['references']}/{r['independent_total']} references, {r['trend']})"
           for r in missing] or ["None."]
    md += ["", f"## Emerging topics not taught ({len(emerging_missing)})", ""]
    md += [f"- **{r['topic']} {r['title']}** ({r['references']} references)" for r in emerging_missing] or ["None."]
    md += ["", "## Taught but weakly supported by the references", ""]
    md += [f"- **{o['topic']} {o['title']}**: {o['hours']} h ({o['share']:.0%} of the course); "
           f"references give it {o['reference_emphasis']:.0%} on average ({o['tier']}, {o['trend']})" for o in overweight] or ["None."]
    md += ["", "## Taught but in no reference", ""]
    md += [f"- Lecture {c['lecture']}: {c['label']} ({c['hours']} h)" for c in current_only] or ["None."]
    md += ["", "## Topics taught in more than one lecture", ""]
    md += [f"- **{o['topic']} {o['title']}**: lectures {', '.join(map(str, o['lectures']))}" for o in overlaps] or ["None."]
    (ws / "gap_analysis.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    return {"ok": True, "out": str(ws / "gap_analysis.md"), "missing": len(missing),
            "emerging_missing": len(emerging_missing), "overweight": len(overweight),
            "current_only": len(current_only), "overlaps": len(overlaps)}


# ---------------------------------------------------------------- field trends

CURRENT = ("growing", "rising in one family", "mixed")
FALLING = ("falling in one family", "declining")
ESTABLISHED = ("generally available", "mature", "established practice")


def _course_attention(linked: list[dict]) -> str:
    """One label for a course topic from the attention verdicts of the trend topics linked to it."""
    v = {x["verdict"] for x in linked}
    rising, falling = v & {"rising in one family", "mixed"}, v & set(FALLING)
    if not linked:
        return "not measured"
    if "growing" in v:
        return "growing"
    if rising:
        return "mixed" if falling else "rising"
    if falling:
        return "steady, partly falling" if "steady" in v else "falling"
    return "steady" if "steady" in v else "too little evidence"


def trends(args) -> dict:
    """Join the field's trend topics (field-trends output) to the course topics through plan/trend_map.json."""
    cfg = _plan(args.plan)
    ws = cfg["_ws"]
    if not cfg.get("trends"):
        raise SgError("plan.yaml names no trends folder (key `trends`)")
    tdir = (cfg["_root"] / cfg["trends"]).resolve()
    attention = load_json(tdir / "attention.json")["topics"]
    profiles = {r["id"]: r for r in load_json(tdir / "profiles.json")["topics"]} if (tdir / "profiles.json").exists() else {}
    taxonomy = {t["id"]: t for t in load_json(ws / "taxonomy.json", "taxonomy")["topics"]}
    tmap = load_json(ws / "trend_map.json", "trend_map")
    new = {n["id"]: n for n in tmap["new_topics"]}
    known = set(taxonomy) | set(new)
    problems = []
    if set(new) & set(taxonomy):
        problems.append(f"new topic ids clash with taxonomy ids: {sorted(set(new) & set(taxonomy))}")
    linked = defaultdict(list)
    seen = set()
    for link in tmap["links"]:
        if link["trend"] not in attention:
            problems.append(f"unknown trend topic {link['trend']}")
            continue
        if link["trend"] in seen:
            problems.append(f"trend topic {link['trend']} is linked twice")
        seen.add(link["trend"])
        if link["topic"] is not None and link["topic"] not in known:
            problems.append(f"{link['trend']} links to unknown course topic {link['topic']}")
        elif link["topic"] is not None:
            a, pr = attention[link["trend"]], profiles.get(link["trend"], {})
            linked[link["topic"]].append({"trend": link["trend"], "name": a["name"], "verdict": a["verdict"],
                                          "maturity": pr.get("status"), "organizations": len(pr.get("organizations", [])),
                                          "adoption": sorted(a.get("adoption", {})),
                                          "mentions": sum(f.get("mentions", 0) for f in a.get("families", {}).values())})
    unlinked = sorted(set(attention) - seen)
    if unlinked:
        problems.append(f"{len(unlinked)} trend topics have no link (topic or null), e.g. {unlinked[:5]}")
    for nid, n in new.items():
        if not linked.get(nid):
            problems.append(f"new topic {nid} has no trend topic linked to it")
        for d in n["depends_on"]:
            if d not in known:
                problems.append(f"new topic {nid} depends on unknown topic {d}")
    out_current = [attention[l["trend"]]["name"] for l in tmap["links"]
                   if l["topic"] is None and l["trend"] in attention and attention[l["trend"]]["verdict"] in CURRENT]
    if problems:
        raise SgError("trend_map.json: " + "; ".join(problems[:12]))
    rows = {}
    for tid in known:
        lk = linked.get(tid, [])
        label = _course_attention(lk)
        mature = [x["maturity"] for x in lk if x["maturity"]]
        rows[tid] = {"attention": label, "linked": lk, "maturity": sorted(set(mature))}
    new_rows = []
    for nid, n in new.items():
        r = rows[nid]
        r["must_decide"] = r["attention"] in ("growing", "rising", "mixed") or any(m in ESTABLISHED for m in r["maturity"])
        new_rows.append({"topic": nid, "title": n["title"], "area": n["area"], "tier": "new",
                         "trend": r["attention"], "depends_on": n["depends_on"], "why": n["why"],
                         "must_decide": r["must_decide"]})
    write_json(ws / "trends.json", {"generated_by": "sg course trends", "trends_folder": str(tdir),
                                    "topics": rows, "new_topics": new_rows, "out_of_scope_current": out_current})
    warnings = []
    if (ws / "gap_analysis.json").exists():
        taught = load_json(ws / "gap_analysis.json")["taught"]
    else:
        taught = {}
        warnings.append("gap_analysis.json is missing, so trend_gaps.md lists every topic as not taught; run "
                        "`sg course gaps`, then `sg course trends` again")
    titles = {**{k: v["title"] for k, v in taxonomy.items()}, **{k: v["title"] for k, v in new.items()}}
    lines = [f"# Trends and the course: {cfg.get('name_en') or cfg.get('name')}", "",
             f"Generated {today()} by `sg course trends` from {tdir}. Attention per course topic summarises the "
             "verdicts of the trend topics linked to it; maturity comes from the topic profiles. Each line names "
             "the linked trend topics that rise or fall and counts the others by verdict, naming the most "
             "discussed steady ones.", ""]
    lines += [f"**Warning:** {w}.\n" for w in warnings]

    def item(tid):
        r = rows[tid]
        if not r["linked"]:
            return f"- **{tid} {titles[tid]}**: {r['attention']}; no linked trend topics"
        moving = [x for x in r["linked"] if x["verdict"] in CURRENT + FALLING]
        parts = [f"{len(r['linked'])} linked trend topic" + ("s" if len(r["linked"]) > 1 else "")]
        if moving:
            parts.append("rising or falling: " + ", ".join(
                f"{x['name']} ({x['verdict']}" + (f", {x['maturity']}" if x["maturity"] else "") + ")" for x in moving))
        others = defaultdict(list)
        for x in r["linked"]:
            if x not in moving:
                others[x["verdict"]].append(x)
        for verdict, xs in sorted(others.items(), key=lambda kv: -len(kv[1])):
            part = f"{verdict}: {len(xs)}"
            if verdict == "steady":
                top = [x["name"] for x in sorted(xs, key=lambda x: -x["mentions"])[:NAMED_STEADY]]
                part += f" ({', '.join(top)}{', …' if len(xs) > NAMED_STEADY else ''})"
            parts.append(part)
        return f"- **{tid} {titles[tid]}**: {r['attention']}; " + "; ".join(parts)
    groups = [("Growing or rising, not taught now", [t for t in rows if rows[t]["attention"] in ("growing", "rising", "mixed")
                                                     and t not in taught]),
              ("New course topics proposed by the mapping", sorted(new)),
              ("Growing or rising, already taught", [t for t in rows if rows[t]["attention"] in ("growing", "rising", "mixed")
                                                     and t in taught]),
              ("Falling, taught now", [t for t in rows if "falling" in rows[t]["attention"] and t in taught])]
    for title, ids in groups:
        lines += [f"## {title} ({len(ids)})", ""] + ([item(t) for t in sorted(ids)] or ["None."]) + [""]
    lines += [f"## Current trends outside the course's subject ({len(out_current)})", "",
              ", ".join(out_current) or "None.", ""]
    (ws / "trend_gaps.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return {"ok": True, "new_topics": len(new), "must_decide": sum(r["must_decide"] for r in new_rows),
            "current_out_of_scope": len(out_current), "out": str(ws / "trend_gaps.md"), "warnings": warnings}


def _rows(ws: Path) -> dict:
    """Course topics: the reference curriculum plus any new topics from the field's trends."""
    rows = {r["topic"]: r for r in load_json(ws / "reference_curriculum.json")["topics"]}
    if (ws / "trends.json").exists():
        t = load_json(ws / "trends.json")
        for r in t["new_topics"]:
            rows[r["topic"]] = r
        for tid, r in rows.items():
            r["attention"] = t["topics"].get(tid, {}).get("attention", "not measured")
            r["maturity"] = t["topics"].get(tid, {}).get("maturity", [])
    return rows


# ---------------------------------------------------------------- proposal

def _check_proposal(cfg: dict) -> tuple[dict, list[str], dict]:
    ws = cfg["_ws"]
    rows = _rows(ws)
    prop = load_json(ws / "proposal.json", "proposal")
    count, hours = cfg["lectures"]["count"], cfg["lectures"]["hours"]
    problems = []
    if len(prop["lectures"]) != count:
        problems.append(f"{len(prop['lectures'])} lectures proposed, the course has {count}")
    if sorted(l["n"] for l in prop["lectures"]) != list(range(1, len(prop["lectures"]) + 1)):
        problems.append("lectures must be numbered 1..N")
    placed: dict[str, int] = {}
    per_lecture, slides, minutes = {}, {}, {}
    slot = hours * 60
    if not slide_range(cfg):
        hint = suggested_slides(cfg)
        problems.append("plan.yaml has no lectures.slides (the expected slide count per lecture, as [min, max])"
                        + (f"; the author's decks suggest {hint}" if hint else ""))
    for lec in prop["lectures"]:
        h = sum(t["hours"] for t in lec["topics"])
        per_lecture[lec["n"]] = round(h, 2)
        if h > hours + 0.01:
            problems.append(f"lecture {lec['n']} needs {h:.2f} h, the slot is {hours} h")
        if h < 0.75 * hours:
            problems.append(f"lecture {lec['n']} fills only {h:.2f} of {hours} h")
        for t in lec["topics"]:
            if t["topic"] not in rows:
                problems.append(f"lecture {lec['n']}: unknown topic {t['topic']}")
            elif t["topic"] in placed and placed[t["topic"]] != lec["n"]:
                problems.append(f"topic {t['topic']} is placed in lectures {placed[t['topic']]} and {lec['n']}; "
                                f"split it into two topics or keep it in one lecture")
            placed.setdefault(t["topic"], lec["n"])
        for r in lec.get("reuse", []):
            if r["deck"] not in {d["id"] for d in cfg.get("decks", [])}:
                problems.append(f"lecture {lec['n']}: reuse names unknown deck {r['deck']}")
        est, rng, pace = slide_estimate(cfg, lec), slide_range(cfg), slides_per_minute(cfg)
        slides[lec["n"]] = est
        if not rng:
            continue
        reused = sum(len(r["slides"]) for r in lec.get("reuse", []))
        new = sum(m["minutes"] for m in lec.get("new_material", []))
        content = material_minutes(cfg, lec)
        minutes[lec["n"]] = round(content)
        if not rng[0] <= est <= rng[1]:
            problems.append(f"lecture {lec['n']} has about {est} slides ({reused} reused + {new:g} min of new "
                            f"material); the author expects {rng[0]}-{rng[1]}, which at the author's pace is "
                            f"{rng[0] / pace:.0f}-{min(rng[1] / pace, slot):.0f} min of material: size "
                            f"new_material in minutes so that it fills the lecture")
        if content > slot + SLOT_SLACK:
            problems.append(f"lecture {lec['n']}: its material takes about {content:.0f} min ({reused} reused "
                            f"slides at the author's pace + {new:g} min of new material), the slot is {slot:g} min")
        if abs(content - h * 60) > MAX_MINUTES_OFF:
            problems.append(f"lecture {lec['n']}: its topics take {h * 60:.0f} min but its material takes about "
                            f"{content:.0f} min ({reused} reused slides + {new:g} min of new material); make them "
                            f"agree within {MAX_MINUTES_OFF} min")
    dropped = {d["topic"] for d in prop["dropped"]}
    assumed = set(prop["assumed_prerequisites"])
    for d in dropped | assumed:
        if d not in rows:
            problems.append(f"dropped or assumed topic {d} is unknown")
        if d in placed:
            problems.append(f"topic {d} is both placed and dropped or assumed")
    for tid, r in rows.items():
        must_decide = r["tier"] in ("core", "common") or r["trend"] == "emerging" or r.get("must_decide")
        if must_decide and tid not in placed and tid not in dropped and tid not in assumed:
            problems.append(f"{r['tier']} topic {tid} ({r['title']}) is neither placed nor dropped with a reason")
    for tid, n in placed.items():
        for dep in rows[tid]["depends_on"]:
            if dep in placed and placed[dep] > n:
                problems.append(f"topic {tid} (lecture {n}) depends on {dep}, which comes later (lecture {placed[dep]})")
            elif dep not in placed and dep not in assumed and dep not in dropped:
                problems.append(f"topic {tid} depends on {dep}, which is neither placed nor assumed")
    if not 1 <= len(prop["decisions"]) <= MAX_DECISIONS:
        problems.append(f"{len(prop['decisions'])} decisions; group related changes into at most {MAX_DECISIONS}")
    ids = [d["id"] for d in prop["decisions"]]
    if len(set(ids)) != len(ids):
        problems.append("duplicate decision ids")
    stats = {"per_lecture_hours": per_lecture, "per_lecture_slides": slides, "per_lecture_minutes": minutes,
             "placed": len(placed), "dropped": len(dropped),
             "assumed": len(assumed), "decisions": len(prop["decisions"]),
             "core_placed": sum(rows[t]["tier"] == "core" for t in placed if t in rows),
             "core_total": sum(r["tier"] == "core" for r in rows.values())}
    return prop, problems, stats


def validate(args) -> dict:
    cfg = _plan(args.plan)
    _, problems, stats = _check_proposal(cfg)
    if problems:
        raise SgError("proposal.json: " + "; ".join(problems[:12]))
    return {"ok": True, **stats}


def report(args) -> dict:
    cfg = _plan(args.plan)
    ws = cfg["_ws"]
    prop, problems, stats = _check_proposal(cfg)
    if problems:
        raise SgError("proposal.json is not valid; run `sg course validate` first: " + problems[0])
    rows = _rows(ws)
    gap = load_json(ws / "gap_analysis.json") if (ws / "gap_analysis.json").exists() else None
    review = load_json(ws / "review.json", "course_review") if (ws / "review.json").exists() else None
    current = {c["n"]: c["title"] for c in cfg.get("current", [])}

    def topic_line(t):
        r = rows[t["topic"]]
        trend = r.get("attention", r["trend"]) if r["tier"] == "new" else r["trend"]
        attention = r.get("attention", "–") if r["tier"] != "new" else trend
        return (f"| {r['topic']} {r['title']} | {t['hours']:.2g} h | {t.get('depth', '')} | {r['tier']} | "
                f"{attention} | {', '.join(r.get('maturity', [])) or '–'} |")

    md = [f"# Curriculum proposal: {cfg.get('name_en') or cfg.get('name')}", "",
          f"Generated {today()} by `sg course report`. Level: {cfg.get('level')}; "
          f"{cfg['lectures']['count']} lectures × {cfg['lectures']['hours']} h.", "",
          f"{stats['core_placed']} of {stats['core_total']} core topics are placed; {stats['dropped']} topics are "
          f"dropped with a reason; {stats['assumed']} are assumed from prerequisites.", "",
          "## Overview", "", "| # | Lecture | Now | Hours | Slides |", "|---|---|---|---|---|"]
    for lec in prop["lectures"]:
        title = lec["title"] + (f" / *{lec['title_local']}*" if lec.get("title_local") else "")
        md.append(f"| {lec['n']} | {title} | {current.get(lec['n'], '–')} | "
                  f"{stats['per_lecture_hours'][lec['n']]:.2g} | about {stats['per_lecture_slides'][lec['n']]} |")
    md += ["", "## Why this structure", "", prop["rationale"], ""]
    for lec in prop["lectures"]:
        md += [f"## Lecture {lec['n']}: {lec['title']}", ""]
        if lec.get("title_local"):
            md += [f"*{lec['title_local']}*", ""]
        md += [lec["summary"], "", "**Objectives**", ""] + [f"- {o}" for o in lec["objectives"]]
        md += ["", "| Topic | Time | Depth | Tier | Attention now | Maturity |", "|---|---|---|---|---|---|"] + [
            topic_line(t) for t in lec["topics"]]
        if lec.get("reuse"):
            md += ["", "**Reusable material**", ""] + [
                f"- {r['deck']}, slides {r['slides'][0]}-{r['slides'][-1]}: {r['note']}" for r in lec["reuse"]]
        if lec.get("new_material"):
            md += ["", "**New material needed**", ""] + [f"- {m['what']} ({m['minutes']:g} min)"
                                                         for m in lec["new_material"]]
        md.append("")
    md += ["## Dropped topics", ""] + [f"- **{d['topic']} {rows[d['topic']]['title']}** "
                                       f"({rows[d['topic']]['tier']}): {d['reason']}" for d in prop["dropped"]] or ["None."]
    md += ["", "## Assumed from prerequisites", ""] + [f"- {t} {rows[t]['title']}" for t in prop["assumed_prerequisites"]] or ["None."]
    if gap:
        md += ["", "## Gaps in the current curriculum that this plan closes", ""]
        md += [f"- {r['topic']} {r['title']} ({r['tier']}): now in lecture {next((l['n'] for l in prop['lectures'] if any(t['topic'] == r['topic'] for t in l['topics'])), '–')}"
               for r in gap["missing"]] or ["None."]
    if review:
        md += ["", "## Reviewer notes", "", review["summary"], ""]
        md += [f"- **{i['severity']}, {i['target']}:** {i['issue']} *Suggestion:* {i['suggestion']}" for i in review["issues"]]
    (ws / "curriculum_proposal.md").write_text("\n".join(md) + "\n", encoding="utf-8")

    dec = [f"# Decisions: {cfg.get('name_en') or cfg.get('name')}", "",
           "Every decision is accepted by default. Untick (`- [ ]`) the ones you reject and add a note "
           "under them; the details are listed for information, not as separate items.", ""]
    for d in prop["decisions"]:
        dec += [f"- [x] **{d['id']} · {d['title']}** (lectures {', '.join(map(str, d['lectures']))})", f"  {d['summary']}", ""]
    (ws / "decisions.md").write_text("\n".join(dec) + "\n", encoding="utf-8")
    return {"ok": True, "proposal": str(ws / "curriculum_proposal.md"), "decisions": str(ws / "decisions.md"), **stats}
