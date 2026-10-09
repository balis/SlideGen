"""Lecture building: a slide-by-slide outline, checked against the course plan before the deck is written.

Layout, relative to the directory of the plan manifest (plan.yaml):
  lectures/L<NN>/outline.json   the outline (schema lecture_outline)
  lectures/L<NN>/deck.json      the deck specification for `sg pptx build`
"""
from __future__ import annotations

from collections import Counter
from pathlib import Path

from .course import _plan, slide_range
from .util import SgError, load_json

SHARE_FACTOR = 1.5   # a topic's share of slides may differ from its share of time by at most this factor


def outline_path(cfg: dict, n: int) -> Path:
    return cfg["_root"] / "lectures" / f"L{n:02d}" / "outline.json"


def _deck_slides(path: Path) -> tuple[int, set[int]]:
    """Number of slides in a deck and the numbers of its hidden slides."""
    from pptx import Presentation
    slides = list(Presentation(str(path)).slides)
    return len(slides), {i for i, sl in enumerate(slides, 1) if sl._element.get("show") == "0"}


def _spec_mismatches(cfg: dict, slides: list[dict], spec_path: Path, warnings: list[str]) -> list[str]:
    """The deck specification must build the outline: same slide count, and template slides copied
    exactly where the outline reuses them."""
    spec = load_json(spec_path, "deck_spec")
    base = (spec_path.resolve().parent / spec["base"]).resolve()
    base_id = next((d["id"] for d in cfg.get("decks", []) if (cfg["_root"] / d["path"]).resolve() == base), None)
    problems = []
    if len(spec["slides"]) != len(slides):
        problems.append(f"the spec has {len(spec['slides'])} slides, the outline {len(slides)}")
    for o, sp in zip(slides, spec["slides"]):
        where, copied, reuse = f"slide {o['n']}", sp.get("copy"), o.get("reuse")
        if o["kind"] == "reuse" and reuse and reuse["deck"] == base_id and copied != reuse["slide"]:
            built = f"copies template slide {copied}" if copied else "builds a new slide"
            problems.append(f"{where}: the outline reuses {base_id} slide {reuse['slide']}, but the spec {built}; "
                            f"update the outline (kind new) or the spec")
        elif o["kind"] == "reuse" and reuse and reuse["deck"] != base_id:
            warnings.append(f"{where}: reused from {reuse['deck']}, which is not the template; the spec must rebuild it")
        elif o["kind"] == "new" and copied:
            problems.append(f"{where}: the outline says new, but the spec copies template slide {copied}; "
                            f"mark it reuse")
    return problems


def outline_check(args) -> dict:
    """Check an outline: slide count in the author's range, every planned topic present, topic shares
    close to the planned time, reused slides that exist, previews that name the right lecture."""
    cfg = _plan(args.plan)
    n = args.lecture
    path = Path(args.outline) if args.outline else outline_path(cfg, n)
    outline = load_json(path, "lecture_outline")
    prop = load_json(cfg["_ws"] / "proposal.json", "proposal")
    lec = next((l for l in prop["lectures"] if l["n"] == n), None)
    if lec is None:
        raise SgError(f"proposal.json has no lecture {n}")
    placed = {t["topic"]: l["n"] for l in prop["lectures"] for t in l["topics"]}
    hours = {t["topic"]: t["hours"] for t in lec["topics"]}
    decks = {d["id"]: cfg["_root"] / d["path"] for d in cfg.get("decks", [])}
    slides = outline["slides"]
    problems, warnings = [], []

    if outline["lecture"] != n:
        problems.append(f"the outline is for lecture {outline['lecture']}, not {n}")
    if [s["n"] for s in slides] != list(range(1, len(slides) + 1)):
        problems.append("slides must be numbered 1..N in order")
    rng = slide_range(cfg)
    if not rng:
        problems.append("plan.yaml has no lectures.slides (the expected slide count per lecture)")
    elif not rng[0] <= len(slides) <= rng[1]:
        problems.append(f"{len(slides)} slides; the author expects {rng[0]}-{rng[1]} for a lecture of "
                        f"{cfg['lectures']['hours'] * 60:g} minutes")

    counts: Counter = Counter()
    sizes: dict[str, tuple[int, set[int]]] = {}
    for s in slides:
        t, where = s["topic"], f"slide {s['n']}"
        if t in hours:
            counts[t] += 1
        elif t is not None and t not in placed:
            problems.append(f"{where}: topic {t} is not placed in any lecture of the plan")
        elif t is not None and s.get("preview_of") != placed[t]:
            problems.append(f"{where}: topic {t} is taught in lecture {placed[t]}; mark the slide "
                            f"`preview_of: {placed[t]}` or give it one of this lecture's topics")
        reuse = s.get("reuse")
        if s["kind"] == "reuse":
            if not reuse:
                problems.append(f"{where}: kind reuse needs `reuse` with deck and slide")
            elif reuse["deck"] not in decks:
                problems.append(f"{where}: unknown deck {reuse['deck']} (plan.yaml decks)")
            elif decks[reuse["deck"]].suffix.lower() == ".pptx" and decks[reuse["deck"]].exists():
                if reuse["deck"] not in sizes:
                    sizes[reuse["deck"]] = _deck_slides(decks[reuse["deck"]])
                total, hidden = sizes[reuse["deck"]]
                if reuse["slide"] > total:
                    problems.append(f"{where}: deck {reuse['deck']} has {total} slides, not {reuse['slide']}")
                elif reuse["slide"] in hidden:
                    warnings.append(f"{where}: {reuse['deck']} slide {reuse['slide']} is hidden in its deck")
        elif reuse:
            problems.append(f"{where}: kind new must not name a reused slide")

    if getattr(args, "spec", None):
        problems += _spec_mismatches(cfg, slides, Path(args.spec), warnings)

    for t in hours:
        if not counts[t]:
            problems.append(f"topic {t} is planned for this lecture but has no slide")
    tagged, total_hours = sum(counts.values()), sum(hours.values())
    shares = {}
    for t, c in counts.items():
        slide_share, time_share = c / tagged, hours[t] / total_hours
        shares[t] = f"{c} slides ({slide_share:.0%}) for {time_share:.0%} of the topic time"
        if not 1 / SHARE_FACTOR <= slide_share / time_share <= SHARE_FACTOR:
            problems.append(f"topic {t}: {shares[t]}; move slides between topics or change the plan's hours")
    if problems:
        raise SgError(f"{path}: " + "; ".join(problems[:12]))
    return {"ok": True, "outline": str(path), "slides": len(slides),
            "reused": sum(s["kind"] == "reuse" for s in slides), "new": sum(s["kind"] == "new" for s in slides),
            "already_checked": [s["n"] for s in slides if s.get("checked")],
            "previews": sum(1 for s in slides if s.get("preview_of")), "topics": shares, "warnings": warnings}
